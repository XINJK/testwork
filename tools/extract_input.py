"""testwork 输入提取器（确定性，不调用 LLM）。

用法:
    python tools/extract_input.py                # 扫描 input/ 下所有支持格式
    python tools/extract_input.py <file> [...]   # 只提取指定文件（相对路径按项目根解析）
    python tools/extract_input.py --force        # 已提取过也重新提取

支持:
    .pdf / .docx / .xlsx / .csv / .html / .htm / .md / .txt / .json / .yaml / .yml
输出:
    input/extracted/<stem>.md，带锚点注释:
    PDF       -> <!-- anchor: p<页码> -->
    DOCX      -> <!-- anchor: P<段落序号> --> ；表格 <!-- anchor: T<表格序号> -->
    XLSX/CSV  -> 表格；引用格式 #<sheet> | "原文摘录"（行序 = 含表头行从 1 计）
    HTML      -> <!-- anchor: B<块序号> -->
    MD/TXT/JSON/YAML -> 原样复制，行锚点 #L<行号>

不支持（会给出提示）:
    .xls / .doc / .ppt / .pptx / 图片；请用 WPS/Office 另存为 .xlsx/.docx，或提供文字版。
"""
from __future__ import annotations

import argparse
import csv as _csv
import datetime as _dt
import io
import sys
import unicodedata
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
INPUT_DIR = ROOT / "input"
OUT_DIR = INPUT_DIR / "extracted"

SUPPORTED = {
    ".pdf", ".docx", ".xlsx", ".csv", ".html", ".htm",
    ".md", ".txt", ".json", ".yaml", ".yml",
}
TEXT_LIKE = {".md", ".txt", ".json", ".yaml", ".yml"}

UNSUPPORTED_HINTS = {
    ".xls": "旧版 .xls 不支持，请用 WPS/Office 另存为 .xlsx 后重试",
    ".doc": "旧版 .doc 不支持，请用 WPS/Office 另存为 .docx 后重试",
    ".ppt": "暂不支持 .ppt/.pptx，可先导出为 PDF/Word，或提供文字版",
    ".pptx": "暂不支持 .ppt/.pptx，可先导出为 PDF/Word，或提供文字版",
    ".png": "图片不支持（无 OCR），请提供文字版或可复制文本的 PDF",
    ".jpg": "图片不支持（无 OCR），请提供文字版或可复制文本的 PDF",
    ".jpeg": "图片不支持（无 OCR），请提供文字版或可复制文本的 PDF",
    ".bmp": "图片不支持（无 OCR），请提供文字版或可复制文本的 PDF",
    ".gif": "图片不支持（无 OCR），请提供文字版或可复制文本的 PDF",
    ".webp": "图片不支持（无 OCR），请提供文字版或可复制文本的 PDF",
}


def _safe_stdout() -> None:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except Exception:
        pass


def _escape_cell(value: str) -> str:
    return value.replace("\r\n", "<br>").replace("\n", "<br>").replace("|", "\\|")


def _stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, _dt.datetime):
        return value.isoformat(sep=" ", timespec="seconds")
    if isinstance(value, _dt.date):
        return value.isoformat()
    return str(value)


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("缺少 pypdf，请先运行: python -m pip install pypdf") from exc
    reader = PdfReader(str(path))
    lines = [f"# {path.stem}", ""]
    total_chars = 0
    for i, page in enumerate(reader.pages, start=1):
        text = (page.extract_text() or "").strip()
        total_chars += len(text)
        lines.append(f"<!-- anchor: p{i} -->")
        if text:
            for para in text.split("\n"):
                para = para.strip()
                if para:
                    lines.append(para)
        lines.append("")
    if total_chars < 20:
        lines.insert(2, "<!-- 警告：未提取到有效文字，可能是扫描件/图片型 PDF；请提供文字版 -->")
    return "\n".join(lines)


def extract_docx(path: Path) -> str:
    try:
        import docx  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("缺少 python-docx，请先运行: python -m pip install python-docx") from exc
    document = docx.Document(str(path))
    lines = [f"# {path.stem}", ""]
    for i, para in enumerate(document.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue
        lines.append(f"<!-- anchor: P{i} -->")
        if para.style is not None and para.style.name.startswith("Heading"):
            lines.append(f"## {text}")
        else:
            lines.append(text)
        lines.append("")
    for ti, table in enumerate(document.tables, start=1):
        lines.append(f"<!-- anchor: T{ti} -->")
        for ri, row in enumerate(table.rows):
            cells = [_escape_cell(c.text.strip()) for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
            if ri == 0:
                lines.append("|" + "|".join(["---"] * max(len(cells), 1)) + "|")
        lines.append("")
    return "\n".join(lines)


def extract_xlsx(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("缺少 openpyxl，请先运行: python -m pip install openpyxl") from exc
    wb = load_workbook(str(path), data_only=True, read_only=True)
    lines = [f"# {path.stem}", ""]
    try:
        for ws in wb.worksheets:
            lines.append(f"## {ws.title}")
            lines.append(f'<!-- 引用格式：#{ws.title} | "原文摘录"（行序 = 含表头行从 1 计） -->')
            emitted = 0
            for row in ws.iter_rows(values_only=True):
                values = [_stringify(v) for v in row]
                if not any(v != "" for v in values):
                    continue
                lines.append("| " + " | ".join(_escape_cell(v) for v in values) + " |")
                if emitted == 0:
                    lines.append("|" + "|".join(["---"] * max(len(values), 1)) + "|")
                emitted += 1
            if emitted == 0:
                lines.append("（空工作表）")
            lines.append("")
    finally:
        wb.close()
    return "\n".join(lines)


def extract_csv(path: Path) -> str:
    text = None
    for enc in ("utf-8-sig", "gbk", "utf-8"):
        try:
            text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise SystemExit("无法识别 CSV 文件编码（尝试了 utf-8-sig / gbk / utf-8）")
    rows = [row for row in _csv.reader(io.StringIO(text))]
    lines = [f"# {path.stem}", ""]
    lines.append(f'<!-- 引用格式：#r<行号> | "原文摘录"（行序 = 含表头行从 1 计） -->')
    for i, row in enumerate(rows, start=1):
        lines.append("| " + " | ".join(_escape_cell(c) for c in row) + " |")
        if i == 1:
            lines.append("|" + "|".join(["---"] * max(len(row), 1)) + "|")
    return "\n".join(lines)


def extract_html(path: Path) -> str:
    try:
        import lxml.html as LH  # type: ignore
    except ImportError as exc:  # pragma: no cover
        raise SystemExit("缺少 lxml，请先运行: python -m pip install lxml") from exc
    import re as _re

    raw = path.read_bytes()
    head = raw[:4096]
    charset = None
    m = _re.search(rb"charset\s*=\s*[\"']?([A-Za-z0-9_\-]+)", head, _re.IGNORECASE)
    if m:
        charset = m.group(1).decode("ascii", errors="ignore")
    text: str | None = None
    for enc in filter(None, (charset, "utf-8-sig", "gbk", "utf-8")):
        try:
            text = raw.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    if text is None:
        text = raw.decode("utf-8", errors="replace")

    def norm(text: str) -> str:
        return " ".join(text.split())

    def inside(el, tag: str) -> bool:
        parent = el.getparent()
        while parent is not None:
            if isinstance(parent.tag, str) and parent.tag.lower() == tag:
                return True
            parent = parent.getparent()
        return False

    doc = LH.document_fromstring(text)
    body = doc.find("body")
    root = body if body is not None else doc.getroot()
    blocks: list[tuple[str, int, object]] = []
    for el in root.iter():
        tag = el.tag.lower() if isinstance(el.tag, str) else ""
        if tag in {"h1", "h2", "h3", "h4", "h5", "h6"}:
            text = norm(el.text_content())
            if text:
                blocks.append(("h", int(tag[1]), text))
        elif tag == "p":
            if inside(el, "table") or inside(el, "li"):
                continue
            text = norm(el.text_content())
            if text:
                blocks.append(("p", 0, text))
        elif tag == "li":
            if inside(el, "table") or inside(el, "li"):
                continue
            text = norm(el.text_content())
            if text:
                blocks.append(("li", 0, text))
        elif tag == "table":
            rows: list[list[str]] = []
            for tr in el.iter("tr"):
                cells = [norm(td.text_content()) for td in tr.iter("td", "th")]
                if any(cells):
                    rows.append(cells)
            if rows:
                blocks.append(("table", 0, rows))
        elif tag == "pre":
            text = el.text_content().rstrip()
            if text:
                blocks.append(("pre", 0, text))

    lines = [f"# {path.stem}", ""]
    n = 0
    for kind, level, payload in blocks:
        n += 1
        lines.append(f"<!-- anchor: B{n} -->")
        if kind == "h":
            lines.append("#" * min(level + 1, 4) + " " + str(payload))
        elif kind == "li":
            lines.append(f"- {payload}")
        elif kind == "pre":
            lines.append("```")
            lines.append(str(payload))
            lines.append("```")
        elif kind == "table":
            rows = payload  # type: ignore[assignment]
            for ri, row in enumerate(rows):
                lines.append("| " + " | ".join(_escape_cell(c) for c in row) + " |")
                if ri == 0:
                    lines.append("|" + "|".join(["---"] * max(len(row), 1)) + "|")
        else:
            lines.append(str(payload))
        lines.append("")
    return "\n".join(lines)


def extract_text(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def resolve_out(path: Path) -> Path:
    """同名不同格式时避免互相覆盖：优先 <stem>.md，冲突则退让为 <stem>-<ext>.md。"""
    tags = [path.suffix.lower().lstrip("."), "2", "3", "4"]
    candidates = [OUT_DIR / f"{path.stem}.md"] + [
        OUT_DIR / f"{path.stem}-{tag}.md" for tag in tags
    ]
    for cand in candidates:
        if not cand.exists():
            return cand
        try:
            header = cand.read_text(encoding="utf-8", errors="replace")[:300]
        except OSError:
            header = ""
        if f"<!-- source: {path.name} " in header:
            return cand
    ext = path.suffix.lower().lstrip(".")
    cand = OUT_DIR / f"{path.stem}-{ext}.md"
    i = 2
    while cand.exists():
        try:
            header = cand.read_text(encoding="utf-8", errors="replace")[:300]
        except OSError:
            header = ""
        if f"<!-- source: {path.name} " in header:
            return cand
        cand = OUT_DIR / f"{path.stem}-{ext}-{i}.md"
        i += 1
    return cand


def collect_inputs(explicit: list[str]) -> list[Path]:
    if explicit:
        files: list[Path] = []
        for item in explicit:
            p = Path(item)
            if not p.is_absolute():
                p = (ROOT / item).resolve()
            files.append(p)
        return files
    files: list[Path] = []
    for p in sorted(INPUT_DIR.iterdir()):
        if p.is_file():
            files.append(p)
    return files


def main() -> int:
    _safe_stdout()
    parser = argparse.ArgumentParser(description="testwork 输入提取器")
    parser.add_argument("files", nargs="*", help="要提取的文件；缺省扫描 input/")
    parser.add_argument("--force", action="store_true", help="覆盖已有提取结果")
    args = parser.parse_args()

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    files = collect_inputs(args.files)
    if not files:
        print(f"未找到可提取文件（目录: {INPUT_DIR}）。")
        return 1

    rc = 0
    for path in files:
        if not path.exists():
            print(f"[跳过] 文件不存在: {path}")
            rc = 1
            continue
        suffix = path.suffix.lower()
        if suffix in UNSUPPORTED_HINTS:
            print(f"[不支持] {path.name}：{UNSUPPORTED_HINTS[suffix]}")
            rc = 1
            continue
        if suffix not in SUPPORTED:
            print(f"[跳过] 不支持的格式: {path.name}")
            continue
        out = resolve_out(path)
        if out.exists() and not args.force:
            try:
                head = out.read_text(encoding="utf-8", errors="replace")[:300]
            except OSError:
                head = ""
            if f"<!-- source: {path.name} " in head:
                print(f"[跳过] 已有提取结果（用 --force 覆盖）: {out.relative_to(ROOT)}")
                continue
        try:
            if suffix == ".pdf":
                content = extract_pdf(path)
            elif suffix == ".docx":
                content = extract_docx(path)
            elif suffix == ".xlsx":
                content = extract_xlsx(path)
            elif suffix == ".csv":
                content = extract_csv(path)
            elif suffix in {".html", ".htm"}:
                content = extract_html(path)
            else:
                content = extract_text(path)
        except Exception as exc:
            print(f"[失败] {path.name}: {exc}")
            rc = 1
            continue
        header = (
            f"<!-- source: {path.name} | extracted: "
            f"{_dt.datetime.now().isoformat(timespec='seconds')} -->\n"
        )
        content = unicodedata.normalize("NFKC", content)
        out.write_text(header + content + "\n", encoding="utf-8")
        print(f"[完成] {path.name} -> {out.relative_to(ROOT)}")
        if suffix == ".pdf" and "警告：未提取到有效文字" in content:
            print(f"[警告] {path.name} 可能是扫描件/图片型 PDF，未提取到有效文字。")
    return rc


if __name__ == "__main__":
    raise SystemExit(main())
