"""testwork 用例导出器: Markdown 的 `## 用例清单` 表格 -> Excel(.xlsx)。

用法:
    python tools/export_xlsx.py <case.md> [case2.md ...] [--out <out.xlsx>]

- 每个 md 一个工作表（表名 = 文件名，最多 31 字符）。
- 单元格中的 <br> 会被还原为 Excel 换行，并启用自动换行。
- 转义竖线 `\\|` 会被还原为 `|`。
- 默认输出到第一个 md 所在目录、同名 .xlsx。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
SECTION = "## 用例清单"
MAX_COL_WIDTH = 60
PIPE_TOKEN = "\u0000PIPE\u0000"


def _split_cells(line: str) -> list[str]:
    inner = line.strip().strip("|")
    inner = inner.replace("\\|", PIPE_TOKEN)
    cells = [c.strip().replace(PIPE_TOKEN, "|") for c in inner.split("|")]
    return cells


def _is_separator(cells: list[str]) -> bool:
    return bool(cells) and all(set(c) <= {"-", ":", " "} for c in cells)


def parse_tables(md_path: Path) -> list[list[list[str]]]:
    """返回文件中所有 markdown 表格（表 = 行列表，行 = 单元列表）。"""
    tables: list[list[list[str]]] = []
    current: list[list[str]] = []
    for raw in md_path.read_text(encoding="utf-8").splitlines():
        line = raw.strip()
        if line.startswith("|") and line.endswith("|"):
            cells = _split_cells(line)
            if _is_separator(cells):
                continue
            current.append(cells)
        else:
            if current:
                tables.append(current)
                current = []
    if current:
        tables.append(current)
    return tables


def load_rows(md_path: Path) -> list[list[str]]:
    """优先取 `## 用例清单` 下的第一张表；缺省取文件最后一张表。"""
    tables = parse_tables(md_path)
    if not tables:
        return []
    text = md_path.read_text(encoding="utf-8")
    if SECTION in text:
        after = text.split(SECTION, 1)[1]
        first_table: list[list[str]] = []
        for raw in after.splitlines():
            line = raw.strip()
            if line.startswith("|") and line.endswith("|"):
                cells = _split_cells(line)
                if _is_separator(cells):
                    continue
                first_table.append(cells)
            elif first_table:
                break
        if first_table:
            return first_table
    return tables[-1]


def sheet_title(md_path: Path) -> str:
    return md_path.stem[:31] or "Sheet1"


def export(paths: list[Path], out: Path) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for md_path in paths:
        rows = load_rows(md_path)
        ws = wb.create_sheet(title=sheet_title(md_path))
        if not rows:
            ws.cell(row=1, column=1, value=f"（{md_path.name} 中未找到可导出的表格）")
            continue
        for r, row in enumerate(rows, start=1):
            for c, value in enumerate(row, start=1):
                cell = ws.cell(row=r, column=c, value=value.replace("<br>", "\n"))
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                if r == 1:
                    cell.font = Font(bold=True)
        for c in range(1, len(rows[0]) + 1):
            width = max(
                (len(str(rows[r][c - 1])) if c - 1 < len(rows[r]) else 0)
                for r in range(len(rows))
            )
            ws.column_dimensions[get_column_letter(c)].width = min(max(width + 4, 10), MAX_COL_WIDTH)
        ws.freeze_panes = "A2"
    if not wb.sheetnames:
        print("没有可导出的内容。")
        return
    wb.save(out)
    print(f"[完成] {out}")


def main() -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8")  # type: ignore[union-attr]
    except Exception:
        pass
    parser = argparse.ArgumentParser(description="testwork 用例 md 导出为 xlsx")
    parser.add_argument("cases", nargs="+", help="用例 md 路径")
    parser.add_argument("--out", default=None, help="输出 xlsx 路径")
    args = parser.parse_args()

    paths: list[Path] = []
    for item in args.cases:
        p = Path(item)
        if not p.is_absolute():
            p = (ROOT / item).resolve()
        if not p.exists():
            print(f"[错误] 文件不存在: {p}")
            return 1
        paths.append(p)
    out = Path(args.out) if args.out else paths[0].with_suffix(".xlsx")
    if not out.is_absolute():
        out = (ROOT / out).resolve()
    export(paths, out)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
